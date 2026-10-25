"""Leer imagenes dentro de PDF y Word (spec leer-imagenes-pdf-word).

Se prueba por la interfaz publica: la tool `read` con el parametro nuevo `images`. El modelo (`complete_with_proveedor`)
y el catalogo son falsos: cero red, cero login del proveedor de IA (ADR-015). Los PDF y el Word con imagenes se GENERAN aca
(pypdf + Pillow + python-docx); el Word trae imagenes de verdad, incluida una dentro de una tabla.

Fase 1: cada caso se escribio como `xfail(strict=True)` y se verifico que fallaba sin la implementacion; en la fase 2 se quitaron.
"""

from __future__ import annotations

import io
import json
import re
from types import SimpleNamespace

import pytest

from app import attachment_readers as ar
from app import main
from app import pdf_vision as pv
from app.tools import filesystem as fs
from test_read_attachment_formats import _pdf_with_text

NOTICE = "(transcripción por visión: puede tener errores)"

CATALOG = {
    "claude-sonnet": {
        "id": "claude-sonnet",
        "supports_image_input": True,
        "context_window": 1_000_000,
        "aliases": ["sonnet"],
    },
    "gpt-5-mini": {"id": "gpt-5-mini", "supports_image_input": True, "context_window": 400_000},
    "gemini-chico": {"id": "gemini-chico", "supports_image_input": True, "context_window": 128_000},
    "deepseek": {"id": "deepseek", "supports_image_input": False, "context_window": 128_000},
}
CATALOG["sonnet"] = CATALOG["claude-sonnet"]  # el catalogo real guarda la misma entrada bajo cada alias


# ---------------------------------------------------------------------------
# Documentos de prueba
# ---------------------------------------------------------------------------
def _png_page(width: int, height: int, color: tuple[int, int, int]):
    import pypdf
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buf, format="PDF")
    return pypdf.PdfReader(io.BytesIO(buf.getvalue())).pages[0]


def build_pdf(pages: list[dict]) -> bytes:
    """Cada pagina: {"text": str, "images": [(ancho_px, alto_px, color), ...]}. El texto es real (extraible)."""
    import pypdf
    from pypdf import Transformation

    writer = pypdf.PdfWriter()
    for spec in pages:
        page = writer.add_blank_page(612, 792)
        if spec.get("text"):
            page.merge_page(pypdf.PdfReader(io.BytesIO(_pdf_with_text([spec["text"]]))).pages[0])
        for i, (w, h, color) in enumerate(spec.get("images", [])):
            page.merge_transformed_page(_png_page(w, h, color), Transformation().translate(50, 640 - i * 55))
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def distinct(n: int, size: int = 100) -> list[tuple[int, int, tuple[int, int, int]]]:
    return [(size, size, (10 + i * 17, 200 - i * 9, 40 + i * 5)) for i in range(n)]


TEXT = "Texto digital de la pagina con contenido suficiente para no parecer escaneada"


def pdf_file(tmp_path, pages: list[dict], name="doc.pdf"):
    f = tmp_path / name
    f.write_bytes(build_pdf(pages))
    return f


def build_docx(path, *, body: str = "full"):
    """Word real: parrafo, imagen, parrafo, tabla con una imagen en una celda, parrafo final. Mas una imagen diminuta (filtrada)
    y la MISMA imagen azul repetida (deduplicada)."""
    import docx
    from docx.shared import Inches
    from PIL import Image

    def png(w, h, color):
        buf = io.BytesIO()
        Image.new("RGB", (w, h), color).save(buf, format="PNG")
        buf.seek(0)
        return buf

    d = docx.Document()
    d.add_paragraph("Antes de la imagen")
    d.add_picture(png(200, 120, (10, 120, 200)), width=Inches(2))
    d.add_paragraph("Despues de la imagen")
    table = d.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.paragraphs[0].add_run("Celda").add_picture(png(100, 100, (200, 0, 0)), width=Inches(1))
    d.add_picture(png(20, 20, (0, 0, 0)), width=Inches(0.3))  # diminuta: no cuenta
    d.add_picture(png(200, 120, (10, 120, 200)), width=Inches(2))  # repetida: no cuenta
    d.add_paragraph("Final del documento")
    d.save(str(path))
    return path


class Model:
    """`complete_with_proveedor` falso: registra cuantas imagenes recibe y devuelve un texto numerado."""

    def __init__(self, fail_on: set[int] | None = None):
        self.calls: list[dict] = []
        self.fail_on = fail_on or set()

    def __call__(self, model, messages):
        content = messages[0]["content"]
        uri = next(b["image_url"]["url"] for b in content if b["type"] == "image_url")
        n = len(self.calls) + 1
        self.calls.append({"model": model, "uri": uri, "prompt": content[0]["text"]})
        if n in self.fail_on:
            raise RuntimeError("falla simulada del modelo")
        return f"TRANSCRIPCION{n}"


@pytest.fixture()
def vision(monkeypatch):
    model = Model()
    monkeypatch.setattr(pv, "complete_with_proveedor", model)
    monkeypatch.setitem(main._MODELS_CACHE, "by_id", CATALOG)
    return model


def session(model="claude-sonnet"):
    return SimpleNamespace(active_model=model, session_id="")


def run_read(path, session_=None, **args):
    return fs._read_execute({"path": str(path), **args}, session_ or session())


# ---------------------------------------------------------------------------
# REQ-003: sin `images`, aviso corto (debajo del umbral)
# ---------------------------------------------------------------------------
def test_below_the_threshold_a_short_notice_is_added_and_nothing_is_read(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(2)}])

    out = run_read(f)

    assert "2 imágenes útiles" in out and "images=true" in out
    assert "3.200 tokens" in out  # 2 x 1.600, es-AR
    assert "question" not in out  # debajo del umbral NO se pregunta
    assert vision.calls == []  # y no se gasta ningun token


def test_without_useful_images_the_output_is_exactly_what_it_was(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT}])

    assert run_read(f) == ar.read_pdf(f)
    assert "imágenes" not in run_read(f)


def test_the_notice_is_appended_after_the_unchanged_text(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(2)}])

    out = run_read(f)

    assert out.startswith(ar.read_pdf(f))  # el texto de siempre, intacto
    assert len(out) > len(ar.read_pdf(f))


# ---------------------------------------------------------------------------
# REQ-002: filtro (tamano y repetidas)
# ---------------------------------------------------------------------------
def test_tiny_and_repeated_images_do_not_count(tmp_path, vision):
    same = (150, 100, (10, 120, 200))
    f = pdf_file(
        tmp_path,
        [{"text": TEXT, "images": [(2, 2, (0, 0, 0)), (8, 8, (1, 1, 1)), same, same, (63, 200, (5, 5, 5))]}],
    )

    out = run_read(f)

    assert "1 imagen útil" in out  # solo `same` (una vez); 2x2, 8x8 y 63 px de lado quedan fuera
    assert "5 imágenes" not in out and "2 imágenes útiles" not in out


def test_the_minimum_size_is_configurable(tmp_path, vision):
    from app import project_context as pc

    f = pdf_file(tmp_path, [{"text": TEXT, "images": [(150, 100, (10, 120, 200)), (300, 300, (9, 9, 9))]}])
    pc.set_image_thresholds(min_px=200)

    out = run_read(f)

    assert "1 imagen útil" in out  # 150x100 queda fuera, 300x300 entra


# ---------------------------------------------------------------------------
# REQ-004: costo alto -> instruccion para preguntar con `question`
# ---------------------------------------------------------------------------
def test_above_the_count_threshold_the_model_is_told_to_ask_with_question(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(6)}])

    out = run_read(f)

    assert "question" in out and "6 imágenes" in out
    assert "9.600 tokens" in out
    assert "Sí, leer las imágenes" in out and "No, solo el texto" in out and "Cambiar de modelo" in out
    assert "(Recomendada)" not in out  # decision del usuario: lo decide una persona, no el /loop
    assert vision.calls == []


def test_above_the_context_percentage_alone_the_model_is_told_to_ask(tmp_path, vision, monkeypatch):
    """3 imagenes (4.800 tokens) estan por debajo de las 5 imagenes, pero son el 24 % de una ventana de 20.000."""
    chico = {"id": "chico", "supports_image_input": True, "context_window": 20_000}
    monkeypatch.setitem(main._MODELS_CACHE, "by_id", {"chico": chico})
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(3)}])

    out = run_read(f, session("chico"))

    assert "question" in out and "24" in out  # el porcentaje se muestra


def test_the_question_shows_the_percentage_of_the_context_window(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(6)}])

    out = run_read(f)  # 9.600 tokens sobre 1.000.000 = 0,96 %

    header = out.split("Opciones:")[0]  # el encabezado de la pregunta; la lista de modelos repite porcentajes y no cuenta
    assert "el 0,96 % de la ventana de contexto del modelo activo (claude-sonnet)" in header


def test_change_model_lists_only_vision_models_largest_window_first_without_inventing_data(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(6)}])

    out = run_read(f)

    assert "deepseek" not in out  # sin vision: no se ofrece
    listing = out[out.rindex("Cambiar de modelo"):]  # la lista, no las menciones anteriores del modelo activo
    assert listing.index("claude-sonnet") < listing.index("gpt-5-mini") < listing.index("gemini-chico")
    assert listing.count("claude-sonnet (") == 1  # el id y su alias son la MISMA entrada del catalogo: se lista una vez
    assert re.findall(r"(?<![\w-])sonnet(?![\w-])", out) == []  # el alias `sonnet` no se lista como un modelo aparte
    low = out.lower()
    assert "ilimitad" not in low and "precio" not in low and "usd" not in low and "$" not in out


# ---------------------------------------------------------------------------
# REQ-001/006: images=true lee; explicito o ya confirmado no vuelve a preguntar
# ---------------------------------------------------------------------------
def test_images_on_a_page_are_read_from_top_to_bottom(tmp_path, vision):
    """Orden de lectura: la imagen 1 es la de mas arriba. En el PDF el eje Y crece hacia arriba, asi que ordenar al reves
    leeria la pagina de abajo hacia arriba. Se identifica cada imagen por su color (rojo arriba, azul abajo)."""
    import base64

    from PIL import Image

    f = pdf_file(tmp_path, [{"text": TEXT, "images": [(100, 100, (250, 10, 10)), (100, 100, (10, 10, 250))]}])

    run_read(f, images=True)

    def center(call):
        img = Image.open(io.BytesIO(base64.b64decode(call["uri"].split(",", 1)[1]))).convert("RGB")
        return img.getpixel((img.width // 2, img.height // 2))

    first, second = (center(c) for c in vision.calls)
    assert first[0] > first[2] and second[2] > second[0]  # rojo primero, azul despues


def test_images_true_reads_the_images_in_order_with_page_and_notice(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(2)}, {"text": TEXT + " dos", "images": distinct(1, 120)}])

    out = run_read(f, images=True)

    assert "[Imagen 1, página 1: TRANSCRIPCION1]" in out
    assert "[Imagen 2, página 1: TRANSCRIPCION2]" in out
    assert "[Imagen 3, página 2: TRANSCRIPCION3]" in out
    assert out.count(NOTICE) == 1
    assert (
        out.index("--- Página 1 ---") < out.index("[Imagen 1") < out.index("--- Página 2 ---") < out.index("[Imagen 3")
    )
    assert len(vision.calls) == 3 and all(c["model"] == "claude-sonnet" for c in vision.calls)


def test_explicit_images_true_with_a_high_cost_reads_without_asking_again(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(8)}])

    out = run_read(f, images=True)

    assert "question" not in out and "[Imagen 8, página 1:" in out
    assert len(vision.calls) == 8


def test_a_confirmed_file_is_not_asked_about_again_in_the_same_session(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(6)}])
    s = session()

    assert "question" in run_read(f, s)  # primera vez, sin images: pregunta
    run_read(f, s, images=True)  # la persona dijo que si
    again = run_read(f, s)  # otra lectura del mismo archivo en la misma sesion

    assert "question" not in again
    assert "6 imágenes útiles" in again  # igual queda el aviso corto: leer sigue pidiendo images=true


def test_the_confirmation_of_one_file_does_not_cover_another(tmp_path, vision):
    a = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(6)}], "a.pdf")
    b = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(7)}], "b.pdf")
    s = session()
    run_read(a, s, images=True)

    assert "question" in run_read(b, s)


# ---------------------------------------------------------------------------
# REQ-005: modelo sin vision
# ---------------------------------------------------------------------------
def test_a_model_without_vision_is_blocked_and_nothing_is_invented(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(2)}])

    out = run_read(f, session("deepseek"), images=True)

    assert "no soporta imágenes" in out
    assert "TRANSCRIPCION" not in out and "[Imagen" not in out
    assert vision.calls == []


def test_a_model_without_vision_is_not_asked_about_a_costly_read(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(8)}])

    out = run_read(f, session("deepseek"))

    assert "question" not in out
    assert "no soporta imágenes" in out or "no ve imágenes" in out


# ---------------------------------------------------------------------------
# REQ-009: tope de 10, imagen que falla, tamano enviado
# ---------------------------------------------------------------------------
def test_at_most_ten_images_per_call_and_the_rest_are_announced(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(12)}])

    out = run_read(f, images=True)

    assert len(vision.calls) == 10 and "[Imagen 10," in out and "[Imagen 11," not in out
    assert "2 imágenes más" in out and "pages" in out


def test_a_failing_image_is_reported_and_the_others_survive(tmp_path, monkeypatch):
    model = Model(fail_on={2})
    monkeypatch.setattr(pv, "complete_with_proveedor", model)
    monkeypatch.setitem(main._MODELS_CACHE, "by_id", CATALOG)
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(3)}])

    out = run_read(f, images=True)

    assert "[Imagen 1, página 1: TRANSCRIPCION1]" in out and "[Imagen 3, página 1: TRANSCRIPCION3]" in out
    assert "TRANSCRIPCION2" not in out
    assert "No pude transcribir la imagen 2" in out  # se informa, no se inventa


def test_the_image_sent_to_the_model_is_a_jpeg_capped_at_1568_px(tmp_path, vision):
    import base64

    from PIL import Image

    f = pdf_file(tmp_path, [{"text": TEXT, "images": [(3000, 2000, (30, 60, 90))]}])

    run_read(f, images=True)

    uri = vision.calls[0]["uri"]
    assert uri.startswith("data:image/jpeg;base64,")
    img = Image.open(io.BytesIO(base64.b64decode(uri.split(",", 1)[1])))
    assert img.format == "JPEG" and pv.MAX_SIDE_PX - 2 <= max(img.size) <= pv.MAX_SIDE_PX  # mismo margen que el render de paginas
    assert max(img.size) < 3000  # y de verdad se achico


def test_the_image_prompt_forbids_inventing_content(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(1)}])

    run_read(f, images=True)

    assert "No inventes" in vision.calls[0]["prompt"]


# ---------------------------------------------------------------------------
# Paginas pedidas y paginas escaneadas
# ---------------------------------------------------------------------------
def test_only_the_requested_pages_count(tmp_path, vision):
    f = pdf_file(tmp_path, [{"text": TEXT, "images": distinct(1)}, {"text": TEXT + " b", "images": distinct(3, 130)}])

    out = run_read(f, pages="1")

    assert "1 imagen útil" in out and "3 imágenes" not in out


def test_a_scanned_page_is_transcribed_once_and_its_embedded_image_is_not_read_again(tmp_path, vision, monkeypatch):
    """Una pagina sin texto ya se renderiza ENTERA: leer ademas su imagen incrustada seria pagar dos veces lo mismo."""
    monkeypatch.setattr(pv, "render_page_jpeg", lambda p, n: b"jpeg")
    f = pdf_file(tmp_path, [{"text": "", "images": distinct(1, 300)}])

    out = run_read(f, images=True)

    assert len(vision.calls) == 1  # la pagina, no la pagina + la imagen
    assert "[Imagen" not in out


# ---------------------------------------------------------------------------
# Word (REQ-008): imagenes de verdad, en su lugar del texto
# ---------------------------------------------------------------------------
def test_word_images_are_read_in_their_place_in_the_text_without_a_page(tmp_path, vision):
    f = build_docx(tmp_path / "informe.docx")

    out = run_read(f, images=True)

    assert out.index("Antes de la imagen") < out.index("[Imagen 1: TRANSCRIPCION1]") < out.index("Despues de la imagen")
    assert (
        out.index("Celda") < out.index("[Imagen 2: TRANSCRIPCION2]") < out.index("Final del documento")
    )  # la de la tabla
    assert "página" not in out.split("Word:", 1)[1].split("(transcripción")[0].replace("Páginas", "")  # sin pagina
    assert out.count(NOTICE) == 1 and len(vision.calls) == 2  # la diminuta y la repetida no se leyeron


def test_word_without_images_flag_gets_the_short_notice_and_the_same_text(tmp_path, vision):
    f = build_docx(tmp_path / "informe.docx")

    out = run_read(f)

    assert out.startswith(ar.read_docx(f))
    assert "2 imágenes útiles" in out and "images=true" in out and vision.calls == []


def test_word_above_the_threshold_asks_first(tmp_path, vision):
    import docx
    from docx.shared import Inches

    d = docx.Document()
    for i, (w, h, c) in enumerate(distinct(6)):
        from PIL import Image

        buf = io.BytesIO()
        Image.new("RGB", (w, h), c).save(buf, format="PNG")
        buf.seek(0)
        d.add_paragraph(f"Parrafo {i}")
        d.add_picture(buf, width=Inches(1))
    f = tmp_path / "seis.docx"
    d.save(str(f))

    out = run_read(f)

    assert "question" in out and "6 imágenes" in out and vision.calls == []


def test_word_pages_selects_image_numbers_beyond_the_first_ten(tmp_path, vision):
    import docx
    from docx.shared import Inches
    from PIL import Image

    d = docx.Document()
    for w, h, c in distinct(12):
        buf = io.BytesIO()
        Image.new("RGB", (w, h), c).save(buf, format="PNG")
        buf.seek(0)
        d.add_picture(buf, width=Inches(1))
    f = tmp_path / "doce.docx"
    d.save(str(f))

    first = run_read(f, images=True)
    rest = run_read(f, images=True, pages="11-12")

    assert "[Imagen 10:" in first and "[Imagen 11:" not in first and "2 imágenes más" in first
    assert "[Imagen 11:" in rest and "[Imagen 12:" in rest and "[Imagen 10:" not in rest


# ---------------------------------------------------------------------------
# Politica de costo (funcion pura) y configuracion
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("n", "window", "costly"),
    [
        (5, 1_000_000, False),  # justo en el umbral de cantidad: "mas de 5"
        (6, 1_000_000, True),
        (2, 40_000, False),  # 3.200 tokens = 8 % de 40.000: debajo del 10 %
        (3, 40_000, True),  # 4.800 = 12 %
        (3, 48_000, False),  # 4.800 = EXACTAMENTE 10 %: "mas de 10 %", no "al menos"
        (3, 47_999, True),  # un token menos de ventana y ya pasa del 10 %
        (0, 1_000, False),
    ],
)
def test_cost_policy_boundaries(n, window, costly):
    from app import document_images as di

    assert di.is_costly(n, window, di.ImageThresholds()) is costly


def test_the_image_thresholds_default_and_are_validated():
    from app import project_context as pc

    cfg = pc.get_project_config()
    assert (cfg.image_min_px, cfg.image_ask_count, cfg.image_ask_pct) == (64, 5, 10)
    for bad in ({"min_px": 0}, {"ask_count": 0}, {"ask_pct": 0}, {"ask_pct": 101}):
        with pytest.raises(ValueError, match="."):
            pc.set_image_thresholds(**bad)
    pc.set_image_thresholds(min_px=100, ask_count=3, ask_pct=25)
    cfg = pc.get_project_config()
    assert (cfg.image_min_px, cfg.image_ask_count, cfg.image_ask_pct) == (100, 3, 25)


def test_the_config_endpoint_reads_and_writes_the_thresholds_without_touching_the_attachments_one():
    from fastapi.testclient import TestClient

    client = TestClient(main.app)
    assert client.get("/api/config/document-images").json() == {"min_px": 64, "ask_count": 5, "ask_pct": 10}
    r = client.post("/api/config/document-images", json={"ask_count": 3, "ask_pct": 25})
    assert r.status_code == 200 and r.json() == {"min_px": 64, "ask_count": 3, "ask_pct": 25}
    assert client.post("/api/config/document-images", json={"ask_pct": 500}).status_code == 400
    assert client.get("/api/config/attachments").json() == {"warn_mb": 100, "inplace_mb": 500}  # contrato intacto


def test_the_read_tool_declares_images_and_explains_the_rule():
    props = fs.read_tool.parameters["properties"]
    assert props["images"]["type"] == "boolean"
    text = json.dumps(props["images"], ensure_ascii=False)
    assert "usuario" in text  # solo si lo pidio o lo confirmo
