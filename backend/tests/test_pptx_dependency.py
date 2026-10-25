"""`python-pptx` tiene que venir en TODAS las instalaciones (pedido del usuario, 2026-10-15).

Por que: a un companero tuvieron que instalarsela a mano para leer un .pptx. Dos cosas la dejan de verdad en todos lados y se vigilan aca:
1. `backend/requirements.txt` la declara (de ahi la instalan el entorno de dev y el lanzador del .exe, que reinstala cuando cambia el hash de ese archivo);
2. ese mismo archivo viaja dentro del `.exe` (`build_stage.py` lo copia) y NO esta en `requirements-dev.txt` (si estuviera solo ahi, prod no la tendria).
Ademas se comprueba que se IMPORTA y lee una presentacion real, no solo que figure en un archivo."""

from __future__ import annotations

import io
import re
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parent.parent
REQUIREMENTS = BACKEND / "requirements.txt"
REQUIREMENTS_DEV = BACKEND / "requirements-dev.txt"
ROOT = BACKEND.parent


def _declared(path: Path) -> list[str]:
    lines = path.read_text(encoding="utf-8").splitlines()
    return [
        re.split(r"[<>=!~\[ ;#]", line.strip(), maxsplit=1)[0].lower()
        for line in lines
        if line.strip() and not line.lstrip().startswith("#")
    ]


def test_requirements_declares_python_pptx_with_a_lower_bound():
    text = REQUIREMENTS.read_text(encoding="utf-8")

    assert "python-pptx" in _declared(REQUIREMENTS)
    assert re.search(r"^python-pptx>=\d", text, re.MULTILINE), "tiene que llevar piso de version (>=), como el resto"


def test_python_pptx_is_in_the_runtime_requirements_not_only_in_the_dev_ones():
    """Si solo estuviera en requirements-dev.txt, prod no la tendria: ese archivo nunca viaja."""
    assert "python-pptx" not in _declared(REQUIREMENTS_DEV)


def test_the_runtime_requirements_file_is_what_prod_and_the_exe_ship():
    # ULTRON no tiene el flujo dev->prod de la version de origen (`sync_prod.py`): lo que viaja es lo que copia `build_stage.py`.
    stage = (ROOT / "build_stage.py").read_text(encoding="utf-8")

    assert '("backend/requirements.txt", "backend/requirements.txt")' in stage  # lo que va adentro del .exe


def test_python_pptx_imports_and_reads_a_real_presentation():
    pptx = pytest.importorskip(
        "pptx", reason="python-pptx no esta instalado en este entorno: justo lo que este pedido evita"
    )
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])  # titulo solo
    slide.shapes.title.text = "SPL de prueba"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "Paso 1: escanear la etiqueta"
    buffer = io.BytesIO()
    deck.save(buffer)
    buffer.seek(0)

    texts = [shape.text_frame.text for s in Presentation(buffer).slides for shape in s.shapes if shape.has_text_frame]

    assert texts == ["SPL de prueba", "Paso 1: escanear la etiqueta"]
    assert pptx.__version__  # trae version: es el paquete real
